from pathlib import Path
import re


# =========================================================
# OCR CONFIGURATION
# =========================================================

def configure_tesseract():

    try:

        import pytesseract

    except ImportError:

        raise RuntimeError(
            "pytesseract is not installed. "
            "Run: pip install pytesseract"
        )


    # -----------------------------------------------------
    # Try Tesseract from PATH
    # -----------------------------------------------------

    try:

        pytesseract.get_tesseract_version()

        return pytesseract

    except Exception:

        pass


    # -----------------------------------------------------
    # Common Windows installation paths
    # -----------------------------------------------------

    possible_paths = [

        Path(
            r"C:\Program Files\Tesseract-OCR\tesseract.exe"
        ),

        Path(
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"
        ),

    ]


    for path in possible_paths:

        if path.exists():

            pytesseract.pytesseract.tesseract_cmd = (
                str(path)
            )

            return pytesseract


    raise RuntimeError(
        "Tesseract OCR could not be found. "
        "Install Tesseract OCR on Windows."
    )


# =========================================================
# PDF TEXT EXTRACTION
# =========================================================

def extract_pdf_text(
    file_path: Path,
) -> str:

    from pypdf import PdfReader


    reader = PdfReader(
        str(file_path)
    )


    pages = []


    for page in reader.pages:

        try:

            page_text = (
                page.extract_text()
                or ""
            )

        except Exception:

            page_text = ""


        if page_text.strip():

            pages.append(
                page_text
            )


    return "\n".join(
        pages
    )


# =========================================================
# PDF OCR FALLBACK
# =========================================================

def extract_pdf_with_ocr(
    file_path: Path,
) -> str:

    try:
        import fitz
        import cv2
        import numpy as np

    except ImportError as error:
        raise RuntimeError(
            "OpenCV, NumPy or PyMuPDF is missing for robust PDF OCR. "
            "Run: pip install opencv-python numpy pymupdf"
        ) from error

    pytesseract = configure_tesseract()

    document = fitz.open(
        str(file_path)
    )

    pages = []

    try:

        for page_number in range(
            len(document)
        ):

            page = document[
                page_number
            ]

            pixmap = page.get_pixmap(
                matrix=fitz.Matrix(
                    4,
                    4,
                ),
                alpha=False,
            )

            image = cv2.imdecode(
                np.frombuffer(
                    pixmap.tobytes("png"),
                    dtype=np.uint8,
                ),
                cv2.IMREAD_GRAYSCALE,
            )

            if image is None:
                continue

            # Mild preprocessing only. Heavy thresholding can destroy
            # thin invoice text and digits.
            content_rect = _detect_document_content_rect(
                image
            )

            if content_rect:
                x1, y1, x2, y2 = content_rect
                ocr_image = image[y1:y2, x1:x2]
            else:
                ocr_image = image

            for config in (
                "--oem 3 --psm 11",
                "--oem 3 --psm 6",
            ):

                try:
                    text = pytesseract.image_to_string(
                        ocr_image,
                        config=config,
                    )
                except Exception:
                    text = ""

                if text.strip():
                    pages.append(text)

    finally:
        document.close()

    if not pages:
        return ""

    # Keep the richer OCR page output.
    return "\n".join(
        max(
            pages,
            key=lambda value: len(value.strip()),
        ).splitlines()
    )




# =========================================================
# DOCUMENT CONTENT DETECTION
# =========================================================

def _detect_document_content_rect(
    image,
):

    """Find the dominant non-white document/photo region."""

    try:
        import cv2

    except ImportError:
        return None

    gray = image

    if len(gray.shape) == 3:
        gray = cv2.cvtColor(
            gray,
            cv2.COLOR_BGR2GRAY,
        )

    _, binary = cv2.threshold(
        gray,
        245,
        255,
        cv2.THRESH_BINARY_INV,
    )

    kernel = cv2.getStructuringElement(
        cv2.MORPH_RECT,
        (25, 25),
    )

    binary = cv2.morphologyEx(
        binary,
        cv2.MORPH_CLOSE,
        kernel,
    )

    contours, _ = cv2.findContours(
        binary,
        cv2.RETR_EXTERNAL,
        cv2.CHAIN_APPROX_SIMPLE,
    )

    if not contours:
        return None

    image_height, image_width = gray.shape[:2]

    candidates = []

    for contour in contours:

        x, y, w, h = cv2.boundingRect(
            contour
        )

        area = w * h

        if area < image_width * image_height * 0.20:
            continue

        if w < image_width * 0.60:
            continue

        if h < image_height * 0.60:
            continue

        candidates.append(
            (
                area,
                x,
                y,
                w,
                h,
            )
        )

    if not candidates:
        return None

    candidates.sort(
        reverse=True
    )

    _, x, y, w, h = candidates[0]

    padding = 20

    x1 = max(
        0,
        x - padding,
    )

    y1 = max(
        0,
        y - padding,
    )

    x2 = min(
        image_width,
        x + w + padding,
    )

    y2 = min(
        image_height,
        y + h + padding,
    )

    return x1, y1, x2, y2


# =========================================================
# ROBUST STRUCTURED OCR FOR PHOTO INVOICES
# =========================================================

def _detect_table_geometry(
    image,
):

    """Detect the main invoice table rectangle."""

    try:
        import cv2

    except ImportError:
        return None

    gray = image

    if len(gray.shape) == 3:
        gray = cv2.cvtColor(
            gray,
            cv2.COLOR_BGR2GRAY,
        )

    binary = cv2.adaptiveThreshold(
        gray,
        255,
        cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
        cv2.THRESH_BINARY_INV,
        31,
        15,
    )

    vertical_kernel = cv2.getStructuringElement(
        cv2.MORPH_RECT,
        (1, max(40, gray.shape[0] // 55)),
    )

    horizontal_kernel = cv2.getStructuringElement(
        cv2.MORPH_RECT,
        (max(40, gray.shape[1] // 35), 1),
    )

    vertical = cv2.morphologyEx(
        binary,
        cv2.MORPH_OPEN,
        vertical_kernel,
    )

    horizontal = cv2.morphologyEx(
        binary,
        cv2.MORPH_OPEN,
        horizontal_kernel,
    )

    combined = cv2.bitwise_or(
        vertical,
        horizontal,
    )

    contours, _ = cv2.findContours(
        combined,
        cv2.RETR_EXTERNAL,
        cv2.CHAIN_APPROX_SIMPLE,
    )

    image_height, image_width = gray.shape[:2]

    candidates = []

    for contour in contours:

        x, y, w, h = cv2.boundingRect(
            contour
        )

        if w < image_width * 0.30:
            continue

        if h < image_height * 0.08:
            continue

        if h > image_height * 0.35:
            continue

        page_position = (
            y / max(image_height, 1)
        )

        if page_position < 0.20:
            continue

        if page_position > 0.80:
            continue

        aspect = w / max(h, 1)

        if aspect < 1.4:
            continue

        candidates.append(
            (
                w * h,
                x,
                y,
                w,
                h,
            )
        )

    if not candidates:
        return None

    candidates.sort(
        key=lambda item: item[0],
        reverse=True,
    )

    _, x, y, w, h = candidates[0]

    return x, y, w, h


def _vertical_table_boundaries(
    table,
):

    try:
        import cv2

    except ImportError:
        return []

    gray = table

    if len(gray.shape) == 3:
        gray = cv2.cvtColor(
            gray,
            cv2.COLOR_BGR2GRAY,
        )

    binary = cv2.adaptiveThreshold(
        gray,
        255,
        cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
        cv2.THRESH_BINARY_INV,
        31,
        15,
    )

    kernel = cv2.getStructuringElement(
        cv2.MORPH_RECT,
        (1, max(25, gray.shape[0] // 12)),
    )

    vertical = cv2.morphologyEx(
        binary,
        cv2.MORPH_OPEN,
        kernel,
    )

    strength = (
        vertical > 0
    ).sum(axis=0)

    threshold = max(
        gray.shape[0] * 0.25,
        70,
    )

    raw = [
        index
        for index, value in enumerate(strength)
        if value >= threshold
    ]

    if not raw:
        return []

    groups = []

    for position in raw:

        if (
            not groups
            or position - groups[-1][-1] > 6
        ):
            groups.append([position])
        else:
            groups[-1].append(position)

    centers = [
        int(sum(group) / len(group))
        for group in groups
    ]

    return centers


def _ocr_table_data(
    table,
):

    try:
        import pytesseract

    except ImportError:
        return {}

    try:
        data = pytesseract.image_to_data(
            table,
            config="--oem 3 --psm 11",
            output_type=pytesseract.Output.DICT,
        )
    except Exception:
        return {}

    rows = []

    for index, value in enumerate(
        data.get("text", [])
    ):

        value = (value or "").strip()

        if not value:
            continue

        try:
            confidence = float(
                data["conf"][index]
            )
        except Exception:
            confidence = 0.0

        if confidence < 20:
            continue

        rows.append(
            {
                "text": value,
                "confidence": confidence,
                "left": int(data["left"][index]),
                "top": int(data["top"][index]),
                "width": int(data["width"][index]),
                "height": int(data["height"][index]),
            }
        )

    return rows


def _numeric_token(
    value: str,
) -> float | int | None:

    match = re.search(
        r"(?<![A-Za-z])([0-9][0-9,]*(?:\.[0-9]+)?)",
        value,
    )

    if not match:
        return None

    return clean_number(
        match.group(1)
    )


def _best_column_number(
    tokens,
    left: int,
    right: int,
    y_min: int = 30,
    y_max: int = 180,
):

    candidates = []

    for token in tokens:

        center_x = (
            token["left"]
            + token["width"] / 2
        )

        center_y = (
            token["top"]
            + token["height"] / 2
        )

        if not (
            left <= center_x <= right
            and y_min <= center_y <= y_max
        ):
            continue

        number = _numeric_token(
            token["text"]
        )

        if number is None:
            continue

        candidates.append(
            (
                token["confidence"],
                number,
            )
        )

    if not candidates:
        return None

    candidates.sort(
        key=lambda item: item[0],
        reverse=True,
    )

    return candidates[0][1]


def _description_from_tokens(
    tokens,
    left: int,
    right: int,
):

    candidates = []

    for token in tokens:

        center_x = (
            token["left"]
            + token["width"] / 2
        )

        center_y = (
            token["top"]
            + token["height"] / 2
        )

        if not (
            left <= center_x <= right
            and 30 <= center_y <= 180
        ):
            continue

        if _numeric_token(
            token["text"]
        ) is not None:
            continue

        if re.fullmatch(
            r"description|item|product|service",
            token["text"],
            re.IGNORECASE,
        ):
            continue

        candidates.append(
            token
        )

    if not candidates:
        return None

    # The first line of the description column is the first item row.
    candidates.sort(
        key=lambda item: (
            item["top"],
            item["left"],
        )
    )

    first_y = candidates[0]["top"]

    same_line = [
        token
        for token in candidates
        if abs(token["top"] - first_y) <= 12
    ]

    same_line.sort(
        key=lambda item: item["left"]
    )

    text = " ".join(
        token["text"]
        for token in same_line
    )

    return clean_value(
        text
    )


def extract_structured_ocr_fields(
    file_path: Path,
) -> dict:

    """Extract fields directly from the geometry of a photographed PDF invoice."""

    try:
        import fitz
        import cv2
        import numpy as np

    except ImportError:
        return {}

    try:
        document = fitz.open(
            str(file_path)
        )
    except Exception:
        return {}

    try:

        for page in document:

            pixmap = page.get_pixmap(
                matrix=fitz.Matrix(
                    4,
                    4,
                ),
                alpha=False,
            )

            image = cv2.imdecode(
                np.frombuffer(
                    pixmap.tobytes("png"),
                    dtype=np.uint8,
                ),
                cv2.IMREAD_GRAYSCALE,
            )

            if image is None:
                continue

            # Header OCR
            header_crop = image[
                int(image.shape[0] * 0.14):
                int(image.shape[0] * 0.37),
                int(image.shape[1] * 0.15):
                int(image.shape[1] * 0.90),
            ]

            header_tokens = _ocr_table_data(
                header_crop
            )

            # For header text parsing, OCR the crop directly.
            try:
                import pytesseract

                header_text = pytesseract.image_to_string(
                    header_crop,
                    config="--oem 3 --psm 6",
                )

            except Exception:
                header_text = ""

            result = {}

            order_id = first_match(
                header_text,
                [
                    r"(?:order\s*(?:id|no|number))\s*[:;\-]?\s*([A-Z0-9][A-Z0-9\/_-]{2,})",
                ],
            )

            if order_id:
                result["order_id"] = clean_value(
                    order_id
                )
                result["invoice_number"] = clean_value(
                    order_id
                )

            invoice_date = first_match(
                header_text,
                [
                    r"(?:invoice\s*date|date)\s*[:;\-]?\s*([0-9]{1,2}[\/\-.][0-9]{1,2}[\/\-.][0-9]{2,4})",
                    r"(?:invoice\s*date|date)\s*[:;\-]?\s*([0-9]{1,2}[- ](?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*[- ][0-9]{2,4})",
                    r"(?:invoice\s*date|date)\s*[:;\-]?\s*([0-9]{1,2}[- ](?:January|February|March|April|May|June|July|August|September|October|November|December)[a-z]*[- ][0-9]{2,4})",
                ],
            )

            if invoice_date:
                result["invoice_date"] = clean_value(
                    invoice_date
                )

            # Company name from the header.
            header_lines = [
                clean_value(line)
                for line in header_text.splitlines()
                if line.strip()
            ]

            company_candidates = []

            for line in header_lines:

                lower = line.lower()

                if re.search(
                    r"\b(limited|ltd\.?|pvt\.?|inc\.?|corporation|company|traders|enterprise)\b",
                    lower,
                ):

                    if re.search(
                        r"head office|dhaka office|phone|fax|bangladesh",
                        lower,
                    ):
                        continue

                    company_candidates.append(
                        line
                    )

            if company_candidates:

                company_candidates.sort(
                    key=lambda value: (
                        "international" in value.lower(),
                        -len(value),
                    ),
                    reverse=True,
                )

                selected = company_candidates[0]
                upper = selected.upper()

                # Remove common OCR duplication before the final company name.
                repeated = list(
                    re.finditer(
                        r"\bINTERNATIONAL\b",
                        upper,
                    )
                )

                if len(repeated) >= 2:

                    start = repeated[-1].start()
                    before = upper[:start].strip()
                    words = before.split()

                    if words:
                        selected = (
                            words[-1]
                            + " INTERNATIONAL"
                            + upper[
                                start + len("INTERNATIONAL"):
                            ]
                        )

                result["vendor_name"] = clean_value(
                    selected
                )

            # -------------------------------------------------
            # Table geometry
            # -------------------------------------------------

            table_rect = _detect_table_geometry(
                image
            )

            if table_rect:

                tx, ty, tw, th = table_rect

                table = image[
                    ty:ty + th,
                    tx:tx + tw,
                ]

                boundaries = sorted(
                    set(
                        _vertical_table_boundaries(
                            table
                        )
                    )
                )

                if boundaries:

                    if boundaries[0] > 20:
                        boundaries.insert(
                            0,
                            0,
                        )

                    if boundaries[-1] < tw - 20:
                        boundaries.append(
                            tw,
                        )

                if len(boundaries) >= 5:

                    # For the common five-column invoice:
                    # serial | description | unit price | quantity | total
                    desc_left = boundaries[-5]
                    desc_right = boundaries[-4]
                    unit_left = boundaries[-4]
                    unit_right = boundaries[-3]
                    qty_left = boundaries[-3]
                    qty_right = boundaries[-2]
                    total_left = boundaries[-2]
                    total_right = boundaries[-1]

                    tokens = _ocr_table_data(
                        table
                    )

                    # Direct OCR of the first data-row cell preserves model/spec numbers
                    # such as 27.5 and 2.10 better than token filtering.
                    description_crop = table[
                        25:min(150, th),
                        max(0, desc_left + 5):max(
                            desc_left + 6,
                            desc_right - 5,
                        ),
                    ]

                    try:
                        import pytesseract

                        description_crop = cv2.resize(
                            description_crop,
                            None,
                            fx=3,
                            fy=3,
                            interpolation=cv2.INTER_CUBIC,
                        )

                        direct_description = pytesseract.image_to_string(
                            description_crop,
                            config="--oem 3 --psm 6",
                        ).strip()
                    except Exception:
                        direct_description = ""

                    description = None

                    for candidate_line in direct_description.splitlines():

                        candidate_line = re.sub(
                            r"\s+",
                            " ",
                            candidate_line,
                        ).strip(
                            " -_=|~`'\""
                        )

                        if not candidate_line:
                            continue

                        if re.fullmatch(
                            r"description|item|product|service",
                            candidate_line,
                            re.IGNORECASE,
                        ):
                            continue

                        if re.search(
                            r"head office|dhaka office|phone|fax|bangladesh|grand total|taka only",
                            candidate_line,
                            re.IGNORECASE,
                        ):
                            continue

                        description = candidate_line
                        break

                    if not description:
                        description = _description_from_tokens(
                            tokens,
                            desc_left + 5,
                            desc_right - 5,
                        )

                    if description:
                        result["product_description"] = description

                    unit_price = _best_column_number(
                        tokens,
                        unit_left + 5,
                        unit_right - 5,
                    )

                    quantity = _best_column_number(
                        tokens,
                        qty_left + 5,
                        qty_right - 5,
                    )

                    line_amount = _best_column_number(
                        tokens,
                        total_left + 5,
                        total_right - 5,
                    )

                    if unit_price is not None:
                        result["unit_price"] = unit_price

                    if quantity is not None:
                        result["quantity"] = quantity

                    if line_amount is not None:
                        result["line_amount"] = line_amount

                    # Quantity is frequently missed in photographed invoices.
                    # Recover it from amount / unit price when the relationship is exact.
                    if (
                        result.get("quantity") is None
                        and result.get("line_amount") is not None
                        and result.get("unit_price") is not None
                    ):

                        try:

                            derived_quantity = (
                                float(result["line_amount"])
                                / float(result["unit_price"])
                            )

                            if (
                                derived_quantity > 0
                                and abs(
                                    derived_quantity
                                    - round(derived_quantity)
                                ) < 0.02
                            ):

                                result["quantity"] = int(
                                    round(derived_quantity)
                                )

                        except Exception:
                            pass

            # Currency from invoice wording.
            if re.search(
                r"\btaka\b|\bBDT\b|\bTk\.?\b|৳",
                header_text,
                re.IGNORECASE,
            ):
                result["currency"] = "BDT"

            # We only need the first page containing a plausible invoice table.
            if result:
                return result

    finally:
        document.close()

    return {}


# =========================================================
# IMAGE OCR
# =========================================================

def extract_image_text(
    file_path: Path,
) -> str:

    try:

        from PIL import Image

    except ImportError:

        raise RuntimeError(
            "Pillow is not installed. "
            "Run: pip install pillow"
        )


    pytesseract = configure_tesseract()


    try:

        image = Image.open(
            file_path
        )

    except Exception as error:

        raise RuntimeError(
            f"Unable to open image: {error}"
        )


    try:

        text = pytesseract.image_to_string(
            image,
            config="--psm 6"
        )

    except Exception as error:

        raise RuntimeError(
            f"Image OCR failed: {error}"
        )


    return text


# =========================================================
# EXCEL EXTRACTION
# =========================================================

def extract_excel_text(
    file_path: Path,
    extension: str,
) -> str:

    if extension == ".xls":

        raise RuntimeError(
            "Legacy .xls extraction is not enabled yet. "
            "Please convert the file to .xlsx."
        )


    try:

        from openpyxl import load_workbook

    except ImportError:

        raise RuntimeError(
            "openpyxl is not installed. "
            "Run: pip install openpyxl"
        )


    workbook = load_workbook(
        filename=str(file_path),
        read_only=True,
        data_only=True,
    )


    lines = []


    try:

        for worksheet in workbook.worksheets:

            lines.append(
                f"[Sheet: {worksheet.title}]"
            )


            for row in worksheet.iter_rows(
                values_only=True
            ):

                values = []


                for value in row:

                    if value is None:

                        continue


                    value_text = str(
                        value
                    ).strip()


                    if value_text:

                        values.append(
                            value_text
                        )


                if values:

                    lines.append(
                        " | ".join(
                            values
                        )
                    )

    finally:

        workbook.close()


    return "\n".join(
        lines
    )


# =========================================================
# MAIN FILE TEXT EXTRACTION
# =========================================================

def extract_file_text(
    file_path: Path,
    extension: str,
) -> tuple[str, str]:

    extension = extension.lower()


    # -----------------------------------------------------
    # PDF
    # -----------------------------------------------------

    if extension == ".pdf":

        try:

            pdf_text = extract_pdf_text(
                file_path
            )

        except Exception:

            pdf_text = ""


        # Normal PDF text extraction
        if len(
            pdf_text.strip()
        ) >= 20:

            return (
                pdf_text,
                "pdf_text",
            )


        # OCR fallback
        ocr_text = extract_pdf_with_ocr(
            file_path
        )


        if ocr_text.strip():

            return (
                ocr_text,
                "pdf_ocr",
            )


        return (
            "",
            "pdf_no_text",
        )


    # -----------------------------------------------------
    # EXCEL
    # -----------------------------------------------------

    if extension in {
        ".xls",
        ".xlsx",
    }:

        text = extract_excel_text(
            file_path,
            extension
        )


        return (
            text,
            "excel_cells",
        )


    # -----------------------------------------------------
    # IMAGE
    # -----------------------------------------------------

    if extension in {
        ".png",
        ".jpg",
        ".jpeg",
        ".webp",
    }:

        text = extract_image_text(
            file_path
        )


        return (
            text,
            "image_ocr",
        )


    raise RuntimeError(
        f"Unsupported file type: {extension}"
    )


# =========================================================
# TEXT NORMALIZATION
# =========================================================

def normalize_text(
    text: str,
) -> str:

    text = text.replace(
        "\r\n",
        "\n"
    )


    text = text.replace(
        "\r",
        "\n"
    )


    # Collapse spaces but preserve line breaks.
    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )


    # Remove excessive blank lines.
    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text
    )


    return text.strip()


# =========================================================
# REGEX HELPERS
# =========================================================

def first_match(
    text: str,
    patterns: list[str],
    flags: int = re.IGNORECASE,
):

    for pattern in patterns:

        match = re.search(
            pattern,
            text,
            flags
        )


        if match:

            value = (
                match.group(1)
                .strip()
            )


            if value:

                return value


    return None


def clean_value(
    value: str | None,
):

    if value is None:

        return None


    value = re.sub(
        r"\s+",
        " ",
        value
    ).strip()


    value = value.strip(
        " :|-"
    )


    return (
        value
        or None
    )


def clean_number(
    value: str | None,
):

    if not value:

        return None


    value = (
        value
        .replace(",", "")
        .strip()
    )


    try:

        number = float(
            value
        )

    except ValueError:

        return value


    if number.is_integer():

        return int(
            number
        )


    return number


# =========================================================
# VENDOR EXTRACTION
# =========================================================

def extract_vendor_name(
    text: str,
) -> str | None:

    # -----------------------------------------------------
    # Explicit vendor label
    # -----------------------------------------------------

    labelled_vendor = first_match(
        text,
        [

            r"(?:vendor\s*name|vendor|supplier\s*name|supplier|seller\s*name|seller)\s*[:\-]\s*([^\n]+)",

        ]
    )


    if labelled_vendor:

        return clean_value(
            labelled_vendor
        )


    # -----------------------------------------------------
    # Format:
    #
    # INVOICE
    # ABC Traders Ltd.
    # Invoice No: ...
    # -----------------------------------------------------

    lines = [

        line.strip()

        for line in text.split("\n")

        if line.strip()

    ]


    for index, line in enumerate(
        lines
    ):

        if re.fullmatch(
            r"invoice",
            line,
            re.IGNORECASE
        ):

            for next_line in lines[
                index + 1:
                index + 6
            ]:

                if re.search(
                    r"\binvoice\s*(?:no|number|#)\b",
                    next_line,
                    re.IGNORECASE
                ):

                    break


                if re.search(
                    r"^\s*date\b",
                    next_line,
                    re.IGNORECASE
                ):

                    continue


                if re.search(
                    r"^(description|quantity|qty|unit\s*price|amount)$",
                    next_line,
                    re.IGNORECASE
                ):

                    continue


                if (
                    re.search(
                        r"[A-Za-z]",
                        next_line
                    )
                    and len(next_line) >= 3
                ):

                    return clean_value(
                        next_line
                    )


    return None


# =========================================================
# TABLE VALUE EXTRACTION
# =========================================================

def extract_table_values(
    text: str,
) -> tuple[
    float | int | None,
    float | int | None,
    float | int | None
]:

    """
    Extract:

        quantity
        unit price
        line amount

    Supports both:

    1. Horizontal:

       Description Quantity Unit Price Amount
       Product / Service 10 5,000 50,000

    2. Vertical PDF text extraction:

       Description
       Quantity
       Unit Price
       Amount
       Product / Service
       10
       5,000
       50,000
    """


    lines = [

        line.strip()

        for line in text.split("\n")

        if line.strip()

    ]


    header_index = None


    # =====================================================
    # FORMAT 1: ALL HEADERS ON ONE LINE
    # =====================================================

    for index, line in enumerate(
        lines
    ):

        has_description = re.search(
            r"\bdescription\b",
            line,
            re.IGNORECASE
        )


        has_quantity = re.search(
            r"\bquantity\b|\bqty\b",
            line,
            re.IGNORECASE
        )


        has_unit_price = re.search(
            r"\bunit\s*price\b",
            line,
            re.IGNORECASE
        )


        has_amount = re.search(
            r"\bamount\b",
            line,
            re.IGNORECASE
        )


        if (
            has_description
            and has_quantity
            and has_unit_price
            and has_amount
        ):

            header_index = index

            break


    # =====================================================
    # FORMAT 2: HEADERS ON SEPARATE LINES
    # =====================================================

    if header_index is None:

        for index in range(
            len(lines) - 3
        ):

            line_1 = lines[
                index
            ].lower()

            line_2 = lines[
                index + 1
            ].lower()

            line_3 = lines[
                index + 2
            ].lower()

            line_4 = lines[
                index + 3
            ].lower()


            if (

                re.fullmatch(
                    r"description",
                    line_1,
                    re.IGNORECASE
                )

                and

                re.fullmatch(
                    r"quantity|qty",
                    line_2,
                    re.IGNORECASE
                )

                and

                re.fullmatch(
                    r"unit\s*price",
                    line_3,
                    re.IGNORECASE
                )

                and

                re.fullmatch(
                    r"amount",
                    line_4,
                    re.IGNORECASE
                )

            ):

                header_index = index

                break


    if header_index is None:

        return (
            None,
            None,
            None
        )


    # =====================================================
    # READ DATA AFTER HEADER
    # =====================================================

    numeric_values = []


    for line in lines[
        header_index + 1:
    ]:

        # -------------------------------------------------
        # Stop when financial summary starts.
        # -------------------------------------------------

        if re.match(
            r"^(subtotal|sub\s*total|vat|tax|total|grand\s*total|amount\s*due|verification\s*code)\b",
            line,
            re.IGNORECASE
        ):

            break


        numbers = re.findall(
            r"\d[\d,]*(?:\.\d+)?",
            line
        )


        for number in numbers:

            value = clean_number(
                number
            )


            if value is not None:

                numeric_values.append(
                    value
                )


    # -----------------------------------------------------
    # We need at least:
    #
    # quantity
    # unit price
    # amount
    # -----------------------------------------------------

    if len(
        numeric_values
    ) < 3:

        return (
            None,
            None,
            None
        )


    # The last three values before the summary are
    # quantity, unit price and line amount.

    quantity = numeric_values[
        -3
    ]


    unit_price = numeric_values[
        -2
    ]


    line_amount = numeric_values[
        -1
    ]


    return (
        quantity,
        unit_price,
        line_amount
    )


# =========================================================
# SUMMARY VALUE EXTRACTION
# =========================================================

def extract_summary_value(
    lines: list[str],
    label_pattern: str,
) -> float | int | None:

    """
    Supports both:

        Total: 57,500

    and:

        Total
        57,500
    """


    for index, line in enumerate(
        lines
    ):

        # -------------------------------------------------
        # Label + number on same line
        # -------------------------------------------------

        same_line = re.match(
            rf"^\s*{label_pattern}\s*[:\-]?\s*(?:[$€£₹৳]|Tk\.?|BDT)?\s*([0-9][0-9,]*(?:\.[0-9]+)?)\s*$",
            line,
            re.IGNORECASE
        )


        if same_line:

            return clean_number(
                same_line.group(1)
            )


        # -------------------------------------------------
        # Label only, number on next line
        # -------------------------------------------------

        if re.fullmatch(
            label_pattern,
            line,
            re.IGNORECASE
        ):

            for next_index in range(
                index + 1,
                min(
                    index + 3,
                    len(lines)
                )
            ):

                next_line = lines[
                    next_index
                ]


                number_match = re.fullmatch(
                    r"(?:[$€£₹৳]|Tk\.?|BDT)?\s*([0-9][0-9,]*(?:\.[0-9]+)?)",
                    next_line,
                    re.IGNORECASE
                )


                if number_match:

                    return clean_number(
                        number_match.group(1)
                    )


    return None


# =========================================================
# FIELD EXTRACTION
# =========================================================

def extract_invoice_fields(
    raw_text: str,
) -> dict:

    text = normalize_text(
        raw_text
    )


    lines = [

        line.strip()

        for line in text.split("\n")

        if line.strip()

    ]


    # =====================================================
    # INVOICE NUMBER
    # =====================================================

    invoice_number = first_match(
        text,
        [

            r"(?:invoice\s*(?:no|number|#))\s*[:\-]?\s*([A-Z0-9][A-Z0-9\/_-]{2,})",

            r"(?:inv\s*(?:no|number|#))\s*[:\-]?\s*([A-Z0-9][A-Z0-9\/_-]{2,})",

        ]
    )


    # =====================================================
    # VENDOR
    # =====================================================

    vendor_name = extract_vendor_name(
        text
    )


    # =====================================================
    # DATE
    # =====================================================

    invoice_date = first_match(
        text,
        [

            r"(?:invoice\s*date|date)\s*[:;\-]?\s*([0-9]{1,4}[\/\-\.][0-9]{1,2}[\/\-\.][0-9]{1,4})",
            r"(?:invoice\s*date|date)\s*[:;\-]?\s*([0-9]{1,2}[- ](?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*[- ][0-9]{2,4})",
            r"(?:invoice\s*date|date)\s*[:;\-]?\s*([0-9]{1,2}[- ](?:January|February|March|April|May|June|July|August|September|October|November|December)[a-z]*[- ][0-9]{2,4})",

        ]
    )


    # =====================================================
    # VERIFICATION CODE
    # =====================================================

    verification_code = first_match(
        text,
        [

            r"(?:verification\s*code)\s*[:\-]?\s*([A-Z0-9][A-Z0-9_-]{3,30})",

            r"(?:validation\s*code)\s*[:\-]?\s*([A-Z0-9][A-Z0-9_-]{3,30})",

            r"(?:authentication\s*code)\s*[:\-]?\s*([A-Z0-9][A-Z0-9_-]{3,30})",

            r"(?:auth\s*code)\s*[:\-]?\s*([A-Z0-9][A-Z0-9_-]{3,30})",

        ]
    )


    # =====================================================
    # TAX ID
    # =====================================================

    tax_id = first_match(
        text,
        [

            r"(?:tax\s*id)\s*[:\-]?\s*([A-Z0-9_-]{3,40})",

            r"(?:vat\s*(?:no|number))\s*[:\-]?\s*([A-Z0-9_-]{3,40})",

        ]
    )


    # =====================================================
    # TABLE VALUES
    # =====================================================

    (
        quantity,
        unit_price,
        line_amount,
    ) = extract_table_values(
        text
    )


    # =====================================================
    # QUANTITY FALLBACK
    # =====================================================

    if quantity is None:

        quantity = clean_number(
            first_match(
                text,
                [

                    r"\b(?:quantity|qty)\b\s*[:\-]?\s*([0-9]+(?:\.[0-9]+)?)",

                ]
            )
        )


    # =====================================================
    # UNIT PRICE FALLBACK
    # =====================================================

    if unit_price is None:

        unit_price = clean_number(
            first_match(
                text,
                [

                    r"\b(?:unit\s*price|price\s*per\s*unit)\b\s*[:\-]?\s*(?:[$€£₹৳]|Tk\.?|BDT)?\s*([0-9][0-9,]*(?:\.[0-9]{1,2})?)",

                ]
            )
        )


    # =====================================================
    # SUBTOTAL
    # =====================================================

    subtotal = extract_summary_value(
        lines,
        r"(?:subtotal|sub\s*total)"
    )


    # =====================================================
    # TAX / VAT
    # =====================================================

    tax_amount = extract_summary_value(
        lines,
        r"(?:vat|tax)(?:\s*amount)?"
    )


    # =====================================================
    # TOTAL
    # =====================================================

    total_amount = extract_summary_value(
        lines,
        r"(?:grand\s*total|total\s*amount|amount\s*due|total)"
    )


    # =====================================================
    # EXTRA FALLBACK:
    # CALCULATE TOTAL FROM SUBTOTAL + TAX
    # =====================================================

    if (
        total_amount is None
        and subtotal is not None
        and tax_amount is not None
    ):

        try:

            total_amount = (
                subtotal
                + tax_amount
            )

        except Exception:

            total_amount = None


    # =====================================================
    # LINE AMOUNT FALLBACK
    # =====================================================

    if (
        line_amount is None
        and quantity is not None
        and unit_price is not None
    ):

        try:

            line_amount = (
                quantity
                * unit_price
            )

        except Exception:

            line_amount = None


    # =====================================================
    # CURRENCY
    # =====================================================

    currency = first_match(
        text,
        [

            r"(?:currency)\s*[:\-]?\s*([A-Z]{3})",

        ]
    )


    if not currency:

        upper_text = text.upper()


        if (
            "৳" in text
            or re.search(
                r"\bBDT\b|\bTK\.?\b|\bTAKA\b",
                upper_text
            )
        ):

            currency = "BDT"

        elif "$" in text:

            currency = "USD"

        elif "€" in text:

            currency = "EUR"

        elif "£" in text:

            currency = "GBP"


    # =====================================================
    # RESULT
    # =====================================================

    fields = {

        "invoice_number":
            clean_value(
                invoice_number
            ),

        "order_id":
            clean_value(
                first_match(
                    text,
                    [
                        r"(?:order\s*(?:id|no|number))\s*[:;\-]?\s*([A-Z0-9][A-Z0-9\/_-]{2,})",
                    ]
                )
            ),

        "vendor_name":
            clean_value(
                vendor_name
            ),

        "invoice_date":
            clean_value(
                invoice_date
            ),

        "verification_code":
            clean_value(
                verification_code
            ),

        "tax_id":
            clean_value(
                tax_id
            ),

        "quantity":
            quantity,

        "unit_price":
            unit_price,

        "line_amount":
            line_amount,

        "subtotal":
            subtotal,

        "tax_amount":
            tax_amount,

        "total_amount":
            total_amount,

        "currency":
            clean_value(
                currency
            ),

    }


    return fields


# =========================================================
# MAIN EXTRACTION FUNCTION
# =========================================================

def extract_invoice_data(
    file_path: Path,
    extension: str,
) -> dict:

    raw_text, method = (
        extract_file_text(
            file_path=file_path,
            extension=extension,
        )
    )


    normalized_text = normalize_text(
        raw_text
    )


    fields = extract_invoice_fields(
        normalized_text
    )


    # -----------------------------------------------------
    # STRUCTURED RECOVERY FOR SCANNED / PHOTO PDFs
    # -----------------------------------------------------

    if extension.lower() == ".pdf":

        try:
            structured_fields = extract_structured_ocr_fields(
                file_path
            )
        except Exception:
            structured_fields = {}

        # Structured OCR is authoritative for table values and company/date
        # fields when it found a plausible value.
        for key, value in structured_fields.items():

            if value is None:
                continue

            if key in {
                "product_description",
                "quantity",
                "unit_price",
                "line_amount",
                "vendor_name",
                "invoice_date",
                "currency",
            }:
                fields[key] = value

            elif key in {
                "invoice_number",
                "order_id",
            }:
                fields[key] = value

        # Total-only invoices may not print a separate subtotal/tax block.
        if (
            fields.get("total_amount") is None
            and fields.get("line_amount") is not None
        ):
            fields["total_amount"] = fields[
                "line_amount"
            ]

        # Recover line amount from quantity × unit price when needed.
        if (
            fields.get("line_amount") is None
            and fields.get("quantity") is not None
            and fields.get("unit_price") is not None
        ):
            try:
                fields["line_amount"] = round(
                    float(fields["quantity"])
                    * float(fields["unit_price"]),
                    2,
                )
            except Exception:
                pass

        if (
            fields.get("total_amount") is None
            and fields.get("line_amount") is not None
        ):
            fields["total_amount"] = fields[
                "line_amount"
            ]


    extracted_count = sum(
        1
        for value in fields.values()
        if value is not None
    )


    # =====================================================
    # STATUS
    # =====================================================

    if not normalized_text and extracted_count > 0:

        extraction_status = (
            "extracted"
        )

    elif not normalized_text:

        extraction_status = (
            "no_text_found"
        )

    elif extracted_count == 0:

        extraction_status = (
            "text_extracted_no_known_fields"
        )

    else:

        extraction_status = (
            "extracted"
        )


    return {

        **fields,

        "extraction_status":
            extraction_status,

        "extraction_method":
            method,

        "fields_detected":
            extracted_count,

        "raw_text_length":
            len(normalized_text),

        "raw_text_preview":
            normalized_text[:5000],

    }